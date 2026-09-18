import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
RESULTS = ROOT / "results"
OUTPUT = ROOT / "AI_Assignment_1_Report.docx"


def set_cell_fill(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_fill(cell, "263B5A")
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = str(value)
            if row_index % 2 == 1:
                set_cell_fill(cell, "F3F6FA")
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run("- " + text)


def add_caption(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.italic = True
    run.font.size = Pt(9)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 32, 32)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)


def build_report() -> None:
    metrics = json.loads((RESULTS / "metrics.json").read_text(encoding="utf-8"))
    history = metrics["history"]
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI Assignment 1 Report")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Adapted ConvNeXt for MNIST Digit Recognition")
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(38, 59, 90)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run("Student: ____________________    Student ID: ____________________")
    metadata.paragraph_format.space_after = Pt(16)

    doc.add_heading("Overview", level=1)
    doc.add_paragraph(
        "This assignment uses an AI coding agent to select, adapt, implement, and evaluate a "
        "recent deep convolutional neural network. The selected algorithm is ConvNeXt, adapted "
        "to a very small two-stage model for 28 by 28 grayscale MNIST images. The experiment "
        f"uses {metrics['train_size']:,} training images and {metrics['test_size']:,} test images "
        f"and reaches a best test accuracy of {metrics['best_test_accuracy'] * 100:.2f}%."
    )

    doc.add_heading("1 How I Asked AI Tools to Find the Algorithm", level=1)
    doc.add_paragraph(
        "I asked the AI tool to recommend a recent deep CNN that could demonstrate a modern "
        "computer-vision design without requiring a large dataset or long training time. I gave "
        "three constraints: the dataset should be MNIST, the available GPU memory is 8 GB, and "
        "the final project should remain small and easy to reproduce."
    )
    add_bullet(doc, "Find a recent deep CNN suitable for an MNIST pattern-recognition assignment.")
    add_bullet(doc, "Keep the model small enough for an 8 GB GPU and a short training run.")
    add_bullet(doc, "Retain the defining ConvNeXt ideas while reducing stages, blocks, and channels.")
    doc.add_paragraph(
        "The AI compared possible architectures and proposed ConvNeXt because it is a modern "
        "convolutional architecture with a clear relationship to classical CNNs. It can also be "
        "scaled down cleanly. The complete prompt record is included with the source code."
    )

    doc.add_heading("2 Algorithm Description", level=1)
    doc.add_paragraph(
        "ConvNeXt modernizes a convolutional network using design choices associated with newer "
        "vision architectures. Its main block applies a large-kernel depthwise convolution for "
        "spatial feature extraction, LayerNorm for normalization, two pointwise convolutions for "
        "channel mixing, GELU activation, and a residual connection."
    )
    add_table(
        doc,
        ["Component", "Adaptation for this assignment"],
        [
            ("Input", "One grayscale channel at 28 by 28 pixels"),
            ("Stem", "3 by 3 convolution producing 24 channels"),
            ("Stage 1", "Two ConvNeXt blocks at 24 channels"),
            ("Downsampling", "2 by 2 stride-2 convolution"),
            ("Stage 2", "Two ConvNeXt blocks at 48 channels"),
            ("Classifier", "Global average pooling and a 10-class linear layer"),
        ],
        widths=[2.0, 4.7],
    )
    doc.add_paragraph(
        "The adapted network removes the additional stages and wide channel dimensions used by "
        "full-size ConvNeXt variants. This reduces computation while preserving the characteristic "
        "depthwise convolution, normalization, pointwise expansion, GELU, and residual pathway."
    )

    doc.add_heading("3 How AI Implemented the Algorithm", level=1)
    doc.add_paragraph(
        "The AI coding agent generated a PyTorch program with separate modules for two-dimensional "
        "LayerNorm, the adapted ConvNeXt block, and the complete classifier. It also implemented "
        "deterministic subset selection, training and evaluation loops, metric recording, model "
        "checkpoint saving, learning-curve plotting, and sample prediction visualization."
    )
    add_bullet(doc, "CrossEntropyLoss is used for the ten digit classes.")
    add_bullet(doc, "AdamW performs optimization with weight decay.")
    add_bullet(doc, "CUDA is selected automatically when it is available; otherwise the code uses CPU.")
    add_bullet(doc, "A fixed seed makes the selected subsets and initialization reproducible.")

    doc.add_heading("4 Experiment Settings and Results", level=1)
    add_table(
        doc,
        ["Setting", "Value"],
        [
            ("Dataset", metrics["dataset"]),
            ("Training images", f"{metrics['train_size']:,}"),
            ("Test images", f"{metrics['test_size']:,}"),
            ("Epochs", metrics["epochs"]),
            ("Batch size", metrics["batch_size"]),
            ("Optimizer", "AdamW"),
            ("Learning rate", metrics["learning_rate"]),
            ("Weight decay", metrics["weight_decay"]),
            ("Model parameters", f"{metrics['parameters']:,}"),
            ("Device", metrics["device"]),
            ("Training and evaluation time", f"{metrics['elapsed_seconds']:.1f} seconds"),
        ],
        widths=[3.2, 3.5],
    )
    epoch_rows = []
    for index in range(metrics["epochs"]):
        epoch_rows.append(
            (
                index + 1,
                f"{history['train_loss'][index]:.4f}",
                f"{history['train_accuracy'][index] * 100:.2f}%",
                f"{history['test_loss'][index]:.4f}",
                f"{history['test_accuracy'][index] * 100:.2f}%",
            )
        )
    add_table(
        doc,
        ["Epoch", "Train loss", "Train accuracy", "Test loss", "Test accuracy"],
        epoch_rows,
        widths=[0.8, 1.35, 1.5, 1.35, 1.5],
    )
    doc.add_paragraph(
        f"The final test accuracy is {metrics['final_test_accuracy'] * 100:.2f}%, and the best "
        f"test accuracy during the run is {metrics['best_test_accuracy'] * 100:.2f}%. The results "
        "show that a heavily reduced ConvNeXt-style model can still learn useful digit features "
        "from a limited subset of MNIST in a short experiment."
    )

    curves = RESULTS / "training_curves.png"
    predictions = RESULTS / "sample_predictions.png"
    if curves.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(curves), width=Inches(5.7))
        add_caption(doc, "Figure 1 Training and test loss and accuracy")
    if predictions.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(predictions), width=Inches(5.3))
        add_caption(doc, "Figure 2 Example predictions from the test subset")

    doc.add_heading("5 What I Learned", level=1)
    doc.add_paragraph(
        "I learned that selecting an algorithm is not only about choosing the model with the "
        "highest published accuracy. The architecture must match the dataset, hardware, time, "
        "and assignment scope. ConvNeXt was designed for much larger color images, so using it "
        "responsibly on MNIST required reducing its width and depth while keeping its key ideas."
    )
    doc.add_paragraph(
        "I also learned that AI-generated code still needs experimental verification. Device "
        "selection, data downloading, deterministic subsets, saved metrics, and visual checks of "
        "predictions are all necessary to turn an implementation into a reproducible experiment. "
        "The AI tool accelerated algorithm search, programming, debugging, and report drafting, "
        "while the final evidence came from executing the program."
    )

    doc.add_heading("6 Source Code Webpage", level=1)
    doc.add_paragraph(
        "GitHub repository: https://github.com/CookieWinnie/AIassignment1"
    )
    doc.add_paragraph(
        "The repository contains the training program, dependency list, AI prompt record, README, "
        "saved metrics, figures, model checkpoint, and this report."
    )

    references = doc.add_paragraph()
    references.paragraph_format.space_before = Pt(8)
    references.paragraph_format.space_after = Pt(0)
    heading_run = references.add_run("References\n")
    heading_run.font.bold = True
    heading_run.font.size = Pt(12.5)
    body_run = references.add_run(
        "Liu, Z. et al. (2022). A ConvNet for the 2020s. Proceedings of the IEEE/CVF Conference "
        "on Computer Vision and Pattern Recognition, 11976-11986.\n"
        "LeCun, Y., Cortes, C., and Burges, C. J. C. The MNIST Database of Handwritten Digits."
    )
    body_run.font.size = Pt(9.5)

    doc.core_properties.title = "AI Assignment 1 Report Adapted ConvNeXt for MNIST"
    doc.core_properties.subject = "AI-assisted deep CNN implementation and experiment"
    doc.core_properties.author = "Student"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
